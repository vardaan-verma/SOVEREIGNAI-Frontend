"""
spreadsheet.py — xlsx_read and xlsx_write tools using openpyxl.

100% local — no external calls. Works on Windows without LibreOffice.
"""

from __future__ import annotations

import time
from pathlib import Path
from typing import Any, Dict, List, Optional

from mcp_server import config
from mcp_server.utils.validation import validate_path


class SpreadsheetTool:

    # ── xlsx_read ─────────────────────────────────────────────────────────

    def xlsx_read(
        self,
        file_path: str,
        sheet_name: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        Read an Excel (.xlsx) file.

        Returns: sheets (list of names), data (sheet→rows), headers, formulas
        """
        start = time.time()
        resolved = validate_path(file_path, must_exist=True)

        try:
            import openpyxl  # type: ignore
        except ImportError:
            raise RuntimeError("openpyxl is not installed. Run: pip install openpyxl")

        wb = openpyxl.load_workbook(str(resolved), data_only=False)
        sheet_names = wb.sheetnames

        # Filter to a specific sheet if requested
        if sheet_name:
            if sheet_name not in sheet_names:
                raise ValueError(
                    f"Sheet '{sheet_name}' not found. Available: {sheet_names}"
                )
            target_sheets = [sheet_name]
        else:
            target_sheets = sheet_names

        all_data: Dict[str, List[List[Any]]] = {}
        all_headers: Dict[str, List[str]] = {}
        all_formulas: Dict[str, Dict[str, str]] = {}

        for sname in target_sheets:
            ws = wb[sname]
            rows: List[List[Any]] = []
            formulas: Dict[str, str] = {}

            for row_idx, row in enumerate(ws.iter_rows(), start=1):
                row_data = []
                for cell in row:
                    val = cell.value
                    # Capture formula strings
                    if isinstance(val, str) and val.startswith("="):
                        formulas[cell.coordinate] = val
                    row_data.append(val)
                rows.append(row_data)

            all_data[sname] = rows
            all_headers[sname] = (
                [str(c) for c in rows[0]] if rows else []
            )
            all_formulas[sname] = formulas

        return {
            "sheets": sheet_names,
            "data": all_data,
            "headers": all_headers,
            "formulas": all_formulas,
            "read_time_ms": int((time.time() - start) * 1000),
        }

    # ── xlsx_write ────────────────────────────────────────────────────────

    def xlsx_write(
        self,
        data: Dict[str, Any],
        output_name: str,
        sheets: Optional[List[Dict[str, Any]]] = None,
    ) -> Dict[str, Any]:
        """
        Write data to an Excel file with optional sheet formatting definitions.

        data format: { "Sheet1": [[row1col1, row1col2, ...], ...], ... }
                  or { "Sheet1": [{"col1": val, ...}, ...] }

        sheets (optional): [{ "name": str, "headers": [...], "freeze_first_row": bool }]

        Returns: path, status, file_size_bytes
        """
        start = time.time()

        try:
            import openpyxl  # type: ignore
            from openpyxl.styles import Font, PatternFill, Alignment  # type: ignore
        except ImportError:
            raise RuntimeError("openpyxl is not installed. Run: pip install openpyxl")

        wb = openpyxl.Workbook()
        # Remove the default empty sheet
        wb.remove(wb.active)

        # Build per-sheet config lookup
        sheet_cfg: Dict[str, Dict[str, Any]] = {}
        if sheets:
            for s in sheets:
                sheet_cfg[s.get("name", "")] = s

        for sheet_name, sheet_data in data.items():
            ws = wb.create_sheet(title=str(sheet_name))
            cfg = sheet_cfg.get(sheet_name, {})

            if not sheet_data:
                continue

            # Normalise rows: accept list-of-dicts or list-of-lists
            if isinstance(sheet_data[0], dict):
                headers = list(sheet_data[0].keys())
                rows = [[row.get(h, "") for h in headers] for row in sheet_data]
            else:
                # list-of-lists; optional explicit headers in cfg
                headers = cfg.get("headers", [])
                rows = [list(r) for r in sheet_data]

            # Write header row with bold styling
            if headers:
                ws.append(headers)
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(
                        start_color="2F4F8F",
                        end_color="2F4F8F",
                        fill_type="solid",
                    )
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center")

            for row in rows:
                ws.append(row)

            # Auto-fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    try:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    except Exception:
                        pass
                ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

            if cfg.get("freeze_first_row", True) and headers:
                ws.freeze_panes = "A2"

        # Save
        safe_name = output_name.replace(" ", "_")
        dest = config.OUTPUT_DIR / f"{safe_name}.xlsx"
        config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        wb.save(str(dest))

        return {
            "path": str(dest),
            "status": "created",
            "file_size_bytes": dest.stat().st_size,
            "sheets_written": list(data.keys()),
            "write_time_ms": int((time.time() - start) * 1000),
        }
